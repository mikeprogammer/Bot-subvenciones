"""Consulta expedientes con autenticación manual FNMT y exporta a Excel."""
import argparse
import csv
import html
import io
import json
import os
import re
import time
import unicodedata
from pathlib import Path
from datetime import datetime
from uuid import uuid4
from urllib.parse import parse_qs, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.sync_api import sync_playwright


def normalize(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).casefold().split())


def new_output_path(directory):
    # Fecha local de España; Windows puede necesitar el paquete tzdata.
    from zoneinfo import ZoneInfo
    stamp = datetime.now(ZoneInfo("Europe/Madrid")).strftime("%Y-%m-%d_%H-%M-%S")
    return Path(directory) / f"consulta_{stamp}_{uuid4().hex[:8]}.xlsx"


def browser_options(config):
    configured = config.get("browser_executable")
    candidates = [Path(configured)] if configured else [
        Path(os.environ.get("PROGRAMFILES", "C:/Program Files")) / "BraveSoftware/Brave-Browser/Application/brave.exe",
        Path(os.environ.get("PROGRAMFILES(X86)", "C:/Program Files (x86)")) / "BraveSoftware/Brave-Browser/Application/brave.exe",
        Path(os.environ.get("LOCALAPPDATA", "")) / "BraveSoftware/Brave-Browser/Application/brave.exe",
    ]
    for candidate in candidates:
        if candidate.is_file():
            return {"executable_path": str(candidate.resolve()), "headless": False}
    raise FileNotFoundError("No se encuentra Brave. Configura browser_executable con la ruta a brave.exe")


def identifier(cell):
    if cell.data_type == "f":
        raise ValueError("Los identificadores deben ser texto, no fórmulas")
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    # Conserva ceros cuando el Excel usa un formato numérico como 00000000.
    if isinstance(value, int) and cell.number_format and set(cell.number_format) == {"0"}:
        return str(value).zfill(len(cell.number_format))
    return str(value).strip()


def read_rows(path, config):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book[config["sheet"]] if config.get("sheet") else book.active
        header = next(sheet.iter_rows(min_row=1, max_row=1))
        names = [normalize(c.value) for c in header]
        indices = []
        for key in ("dni_column", "expediente_column"):
            name = normalize(config[key])
            if names.count(name) != 1:
                raise ValueError(f"Columna ausente o duplicada: {config[key]}")
            indices.append(names.index(name))
        rows = []
        for number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
            if all(c.value is None for c in cells):
                continue
            try:
                dni, expediente = (identifier(cells[i]) for i in indices)
                error = "Falta DNI o número de expediente" if not dni or not expediente else ""
            except ValueError as exc:
                dni, expediente, error = "", "", str(exc)
            rows.append((number, dni, expediente, error))
        return rows
    finally:
        book.close()


def sheets_export_url(url):
    parsed = urlparse(url)
    match = re.fullmatch(r"/spreadsheets/d/([A-Za-z0-9_-]+)(?:/.*)?", parsed.path)
    if parsed.scheme != "https" or parsed.netloc != "docs.google.com" or not match:
        raise ValueError("Usa la URL https://docs.google.com/spreadsheets/d/... de la hoja")
    gid = parse_qs(parsed.fragment).get("gid", parse_qs(parsed.query).get("gid", []))
    if not gid or len(gid) != 1 or not gid[0].isdigit():
        raise ValueError("La URL debe incluir #gid=... para identificar la pestaña exacta")
    return f"https://docs.google.com/spreadsheets/d/{match[1]}/export?format=csv&gid={gid[0]}"


def parse_sheets_csv(text, config):
    values = list(csv.reader(io.StringIO(text.lstrip('\ufeff'))))
    return parse_sheets_values(values, config)


def parse_sheets_values(values, config):
    if not values:
        raise ValueError("La pestaña de Sheets está vacía")
    header = [normalize(value) for value in values[0]]
    indices = []
    for key in ("dni_column", "expediente_column"):
        name = normalize(config[key])
        if header.count(name) != 1:
            raise ValueError(f"Columna ausente o duplicada en Sheets: {config[key]}")
        indices.append(header.index(name))
    records = []
    for number, row in enumerate(values[1:], 2):
        if not any(str(value).strip() for value in row):
            continue
        dni, expediente = (str(row[i]).strip() if i < len(row) else "" for i in indices)
        error = "Falta DNI o número de expediente" if not dni or not expediente else ""
        records.append((number, dni, expediente, error))
    return records


def save_results(path, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = "Resultados"
    sheet.append(["DNI", "Nº EXPEDIENTE", "Estado", "Anomalias", "Documentos referenciados", "Texto original", "Error de consulta", "Fila origen"])
    for row in rows:
        sheet.append(row)
        # Guarda todo texto literalmente, aunque comience con '='.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="17365D")
    for column, width in zip("ABCDEFGH", (16, 27, 30, 70, 60, 70, 50, 14)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.stem + ".tmp.xlsx")
    book.save(temporary)
    os.replace(temporary, path)


def authenticated_page(context):
    pages = [p for p in context.pages if not p.is_closed()]
    print("Pestañas abiertas:")
    for index, page in enumerate(pages, 1):
        print(f"  {index}: {page.url}")
    choice = int(input("Número de pestaña con el formulario de consulta: "))
    if not 1 <= choice <= len(pages):
        raise ValueError("Número de pestaña inválido")
    return pages[choice - 1]


def inspect_form(page):
    print("URL del formulario:", page.url)
    print("Controles (sin valores personales):")
    print(json.dumps(page.locator("input, button, select").evaluate_all(
        "els => els.map(e => ({tag:e.tagName,type:e.type,id:e.id,name:e.name,"
        "labels:Array.from(e.labels || []).map(l=>l.textContent.trim()),"
        "text:e.tagName==='BUTTON'?e.textContent.trim():null}))"
    ), ensure_ascii=False, indent=2))
    print("Usa el inspector para localizar resultado, estado y anomalías.")
    page.pause()


def document_references(anomaly, links):
    references = []
    heading = re.match(
        r"^((?:document|documento)\s+(?:normalitza[dt]|normalizado)?\s*\d+[A-Za-z]?)\s*:\s*([^:]+)\s*:",
        anomaly, re.I,
    )
    if heading:
        references.append(f"{heading.group(1).strip()}: {heading.group(2).strip()}")
    # Conserva la frase original: no adivina nombres a partir del motivo del error.
    if not heading:
        for match in re.finditer(r"\b(?:documento?s?|documents?|documentació[n]?|documentaci[oó])\b[^\n;]*", anomaly, re.I):
            references.append(match.group(0).strip().rstrip("."))
    for match in re.finditer(r"[\wÀ-ÿ][\wÀ-ÿ .()-]*\.(?:pdf|docx?|xlsx?|jpe?g|png)\b", anomaly, re.I):
        references.append(match.group(0).strip())
    for link in links:
        label = link.get("text", "").strip()
        href = link.get("href", "")
        if label and normalize(label) in normalize(anomaly):
            references.append(f"{label} ({href})" if href else label)
    return list(dict.fromkeys(references))


def parse_result(text, links=()):
    original = text.strip()
    if not original:
        raise ValueError("El resultado está vacío")
    decoded = html.unescape(original).replace('\xa0', ' ')
    if "no s'ha trobat cap expedient" in normalize(decoded):
        return "Expediente no encontrado", "No disponibles: expediente no encontrado", "", original
    marker = re.search(r"\b(?:anomalies|anomal[ií]as)\s*:", decoded, re.I)
    state_text = decoded[:marker.start()].strip() if marker else decoded
    # Si hay etiqueta explícita, usa su valor sin el resto del bloque.
    state_match = re.search(r"\b(?:estat|estado)\s*:\s*([^\n]+)", state_text, re.I)
    state = state_match.group(1).strip() if state_match else "No indicado en el texto"
    if not state:
        raise ValueError("No se ha identificado el estado; revisar el texto del resultado")
    anomalies = []
    if marker:
        tail = decoded[marker.end():].strip()
        # Guion de viñeta: al principio de línea o separado por espacios.
        # No divide nombres como informe-final.pdf ni números como DOC-123.
        chunks = re.split(r"(?:^|\n)\s*[-–•]\s*|[ \t]+-[ \t]+", tail)
        anomalies = [" ".join(chunk.split()) for chunk in chunks if chunk.strip()]
        if not anomalies:
            raise ValueError("La sección Anomalies está vacía; no se confirma su ausencia")
        if len(anomalies) == 1 and normalize(anomalies[0]).rstrip(".") in {
            "cap", "cap anomalia", "cap anomalies", "sense anomalies", "sin anomalias", "ninguna"
        }:
            anomalies = []
    # La ausencia de la sección se conserva sin afirmar que no hay anomalías.
    anomaly_text = "\n".join(f"{i}. {value}" for i, value in enumerate(anomalies, 1))
    if not marker:
        anomaly_text = "No aparece la sección Anomalies en el bloque leído"
    elif not anomalies:
        anomaly_text = "Sin anomalías (indicado expresamente)"
    documents = []
    for index, anomaly in enumerate(anomalies, 1):
        references = document_references(anomaly, links)
        documents.append(f"{index}. " + ("; ".join(references) if references else "Sin documento identificable en el texto"))
    return state, anomaly_text, "\n".join(documents), original


def visible_result_blocks(page, selector):
    return page.locator(selector).evaluate_all("""els => els
        .filter(e => e.getClientRects().length && getComputedStyle(e).visibility !== 'hidden')
        .filter(e => !els.some(other => other !== e && other.contains(e) && other.getClientRects().length))
        .map(e => ({text: e.innerText.trim(), links: Array.from(e.querySelectorAll('a')).map(a =>
            ({text:a.innerText.trim(), href:a.href}))}))
        .filter(e => e.text)""")


def consult(page, form_url, selectors, dni, expediente, timeout_ms=30000):
    # Conserva la pestaña autenticada. Neteja elimina la respuesta anterior sin
    # volver a navegar por el acceso FNMT para cada fila.
    if page.is_closed():
        raise RuntimeError('La pestaña de consulta se ha cerrado. Vuelve a iniciar el lote y mantén Brave abierto.')
    if page.url.split('?')[0].rstrip('/') != form_url.rstrip('/'):
        page.goto(form_url, wait_until="domcontentloaded")
    else:
        reset = page.get_by_role('button', name=re.compile('Neteja'))
        if reset.count() == 1 and reset.is_visible():
            reset.click()
        elif visible_result_blocks(page, selectors['resultado']):
            page.goto(form_url, wait_until='domcontentloaded')
    page.locator(selectors["titular"]).set_checked(False)
    page.locator(selectors["dni"]).fill(dni)
    page.locator(selectors["expediente"]).fill(expediente)
    before = visible_result_blocks(page, selectors["resultado"])
    page.locator(selectors["buscar"]).click()
    deadline = time.monotonic() + timeout_ms / 1000
    while time.monotonic() < deadline:
        blocks = visible_result_blocks(page, selectors["resultado"])
        if blocks and blocks != before:
            if len(blocks) != 1:
                raise ValueError("Hay varios bloques mt-3: configura un selector exclusivo del resultado")
            return parse_result(blocks[0]["text"], blocks[0]["links"])
        page.wait_for_timeout(200)
    raise TimeoutError("No apareció un resultado nuevo tras buscar")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", default="config.json")
    source = parser.add_mutually_exclusive_group()
    source.add_argument("--input", type=Path)
    source.add_argument("--sheets-url", help="URL de Google Sheets con #gid=...")
    parser.add_argument("--output", type=Path, help="Nombre opcional; por defecto genera un Excel nuevo con fecha")
    parser.add_argument("--inspect", action="store_true")
    parser.add_argument("--check-sheets", action="store_true", help="Probar solo lectura de Sheets vía API")
    args = parser.parse_args()
    config = json.loads(Path(args.config).read_text(encoding="utf-8-sig"))
    args.output = args.output or new_output_path(config.get("output_dir", "outputs/consultas"))
    sheets_url = args.sheets_url or (config.get("sheets_url") if not args.input else None)
    if not args.inspect:
        if not args.input and not sheets_url:
            parser.error("Se requiere --sheets-url URL o sheets_url en config.json (alternativa: --input archivo.xlsx)")
        if args.input and args.input.resolve() == args.output.resolve():
            parser.error("La salida debe ser diferente del Excel de entrada")
        if args.output.exists():
            parser.error("La salida ya existe; elige otro nombre para conservarla")
        missing = [key for key in ("titular", "dni", "expediente", "buscar", "resultado")
                   if not config.get("selectors", {}).get(key)]
        if missing:
            parser.error("Configura los selectores: " + ", ".join(missing))
        if sheets_url:
            sheets_export_url(sheets_url)
        records = read_rows(args.input, config) if args.input else None
        if records == []:
            parser.error("El Excel no contiene filas de datos")
        if sheets_url:
            from sheets_api import read_values
            records = parse_sheets_values(read_values(sheets_url, config, browser_options(config)['executable_path']), config)
            if not records:
                parser.error('La pestaña no contiene filas de datos')
            print(f'Sheets API: {len(records)} filas leídas; la hoja original no se modifica')
        if args.check_sheets:
            if not sheets_url:
                parser.error('--check-sheets requiere Google Sheets como entrada')
            print('Lectura comprobada. No se han consultado expedientes ni creado un Excel.')
            return
    with sync_playwright() as playwright:
        context = playwright.chromium.launch_persistent_context(
            str(Path(config.get("browser_profile", ".brave-bot-profile")).resolve()), **browser_options(config)
        )
        try:
            context.set_default_timeout(config.get("timeout_ms", 30000))
            page = context.pages[0] if context.pages else context.new_page()
            page.goto(config["login_url"], wait_until="domcontentloaded")
            input("En Brave pulsa Entra, selecciona FNMT y abre el formulario. Después pulsa Intro aquí: ")
            page = authenticated_page(context)
            if args.inspect:
                inspect_form(page)
                return
            form_url = config.get("form_url") or page.url
            results = []
            for index, (source_row, dni, expediente, error) in enumerate(records, 1):
                state = anomalies = documents = original = ""
                if not error:
                    try:
                        state, anomalies, documents, original = consult(page, form_url, config["selectors"], dni, expediente, config.get("timeout_ms", 30000))
                    except Exception as exc:
                        # No mezcla errores técnicos con las anomalías del expediente.
                        error = type(exc).__name__ + ": " + str(exc).splitlines()[0][:500]
                results.append((dni, expediente, state, anomalies, documents, original, error, source_row))
                save_results(args.output, results)
                print(f"Fila {source_row}: {'ERROR' if error else 'OK'} ({index}/{len(records)})")
                if error and dni and expediente:
                    answer = input("Consulta fallida. Intro para continuar o 'salir' para detener: ")
                    if answer.strip().casefold() == "salir":
                        break
                time.sleep(max(0, config.get("delay_seconds", 2)))
            print(f"Resultados guardados: {args.output.resolve()}")
        finally:
            context.close()


if __name__ == "__main__":
    main()
