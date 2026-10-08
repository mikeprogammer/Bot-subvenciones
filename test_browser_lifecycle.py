import unittest
from unittest.mock import MagicMock, patch
import bot
from app import friendly_error, App


class BrowserLifecycleTests(unittest.TestCase):
    def test_incomplete_row_is_saved_and_next_valid_row_is_consulted(self):
        import queue
        import threading
        import tempfile
        from contextlib import nullcontext
        from openpyxl import load_workbook
        from pathlib import Path
        app = App.__new__(App)
        app.events = queue.Queue()
        app.stop = threading.Event()
        app.resume = threading.Event()
        app.resume.set()
        context = MagicMock()
        page = MagicMock()
        context.pages = [page]
        with tempfile.TemporaryDirectory() as directory:
            config = {'sheets_url': 'test', 'dni_column': 'DNI', 'expediente_column': 'Nº EXPEDIENTE',
                      'browser_profile': 'test', 'timeout_ms': 1000, 'login_url': 'test',
                      'form_url': 'test', 'selectors': {}, 'output_dir': directory, 'delay_seconds': 0}
            with patch('app.read_values', return_value=[['DNI', 'Nº EXPEDIENTE'], ['A', 'EXP1'], ['', 'EXP2'], ['B', 'EXP3']]), \
                 patch('app.certificate_policy', return_value=nullcontext()), \
                 patch('app.browser_options', return_value={}), \
                 patch('app.sync_playwright') as playwright, \
                 patch('app.find_form', return_value=page), \
                 patch('app.consult', return_value=('Pendent', '', '', 'Estat: Pendent')) as consult:
                playwright.return_value.__enter__.return_value.chromium.launch_persistent_context.return_value = context
                app.worker(config, 'brave.exe', False)
            self.assertEqual(consult.call_count, 2)
            book = load_workbook(next(Path(directory).glob('*.xlsx')))
            self.assertEqual(book.active.max_row, 4)
            self.assertEqual(book.active.cell(3, 7).value, 'Falta DNI')
            self.assertEqual(book.active.cell(4, 1).value, 'B')
            book.close()

    def test_hidden_dni_before_unchecking_titular_does_not_block_form(self):
        page = MagicMock()
        page.is_closed.return_value = False
        page.url = 'https://intranet.caib.es/subvenfront/consulta'
        controls = {key: MagicMock() for key in ('titular', 'dni', 'expediente', 'buscar')}
        for control in controls.values():
            control.count.return_value = 1
            control.is_visible.return_value = True
        controls['dni'].is_visible.return_value = False
        page.locator.side_effect = lambda selector: controls[selector]
        self.assertTrue(bot.form_ready(page, page.url, {key: key for key in controls}))
        controls['dni'].count.return_value = 0
        self.assertTrue(bot.form_ready(page, page.url, {key: key for key in controls}))

    def test_form_detected_by_fields_despite_hash_and_duplicate_tabs(self):
        selectors = {k: k for k in ('titular', 'dni', 'expediente', 'buscar')}
        page = MagicMock()
        page.is_closed.return_value = False
        page.url = 'https://intranet.caib.es/subvenfront/#/consulta?test=1'
        page.locator.return_value.count.return_value = 1
        page.locator.return_value.is_visible.return_value = True
        context = MagicMock()
        context.pages = [page, page]
        self.assertIs(bot.find_form(context, 'https://intranet.caib.es/subvenfront/consulta', selectors), page)
        page.url = 'https://otro.example/subvenfront/consulta'
        self.assertIsNone(bot.find_form(context, 'https://intranet.caib.es/subvenfront/consulta', selectors))

    def test_login_page_does_not_count_as_ready(self):
        page = MagicMock()
        page.is_closed.return_value = False
        page.url = 'https://intranet.caib.es/subvenfront/login'
        page.locator.return_value.count.return_value = 0
        self.assertFalse(bot.form_ready(page, 'https://intranet.caib.es/subvenfront/consulta',
                                       {k: k for k in ('titular', 'dni', 'expediente', 'buscar')}))

    def test_closed_tab_stops_before_navigation(self):
        page = MagicMock()
        page.is_closed.return_value = True
        with self.assertRaisesRegex(RuntimeError, 'se ha cerrado'):
            bot.consult(page, 'https://example.com/form', {}, 'DNI', 'EXP')
        page.goto.assert_not_called()

    def test_authenticated_form_uses_clear_and_keeps_current_tab(self):
        page = MagicMock()
        page.is_closed.return_value = False
        page.url = 'https://example.com/form'
        reset = page.get_by_role.return_value
        reset.count.return_value = 1
        reset.is_visible.return_value = True
        selectors = {k: k for k in ('titular', 'dni', 'expediente', 'buscar', 'resultado')}
        blocks = [{'text': "No s'ha trobat cap expedient amb aquestes dades", 'links': []}]
        with patch('bot.visible_result_blocks', side_effect=[[], blocks]):
            result = bot.consult(page, page.url, selectors, '123', 'EXP')
        reset.click.assert_called_once()
        page.goto.assert_not_called()
        self.assertEqual(result[0], 'Expediente no encontrado')

    def test_target_closed_error_is_understandable(self):
        error = RuntimeError('Target page, context or browser has been closed')
        self.assertIn('Mantén abierta', friendly_error(error))
