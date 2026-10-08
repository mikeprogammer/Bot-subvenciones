import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from bot import parse_result, save_results, changed_result, report_status


class ResultTests(unittest.TestCase):
    def test_simple_status_uses_anomalies_before_favorable_and_respects_pending(self):
        self.assertEqual(report_status('La concessió ha estat favorable', '')[0], 'Favorable')
        self.assertEqual(report_status('Resolución favorable. Anomalies: - Falta firma', '1. Falta firma')[0], 'Anomalías')
        for text in ('', 'Estat: Pendent', 'Pendiente de resolución favorable', 'Resolución desfavorable', 'No es favorable'):
            self.assertEqual(report_status(text, '')[0], 'En espera')
        self.assertEqual(report_status('Estat: Concedit. Anomalies: Cap', '')[0], 'Favorable')
        self.assertIn('en espera', report_status('Pendent', '')[1])

    def test_multiple_result_panels_preserve_status_anomalies_and_document_links(self):
        static = {'text': 'Introdueixi les dades per a realitzar la consulta', 'links': []}
        status = {'text': 'Estat: Pendent', 'links': []}
        anomalies = {'text': 'Anomalies:\n- Document 2: Falta signatura',
                     'links': [{'text': 'Document 2', 'href': 'https://example.com/doc2.pdf'}]}
        text, links = changed_result([static, status, anomalies, status], [static])
        result = parse_result(text, links)
        self.assertEqual(result[0], 'Pendent')
        self.assertIn('Falta signatura', result[1])
        self.assertIn('doc2.pdf', result[2])
        self.assertNotIn(static['text'], result[3])
        self.assertEqual(result[3].count('Estat:'), 1)

    def test_not_found_is_search_result_not_empty_anomalies(self):
        state, anomalies, documents, original = parse_result("No s'ha trobat cap expedient amb aquestes dades")
        self.assertEqual(state, 'Expediente no encontrado')
        self.assertIn('No disponibles', anomalies)
        self.assertEqual(documents, '')

    def test_anomalies_documents_and_hyphens(self):
        text = 'Estat: Pendent\nAnomalies:\n- Falta signatura al document informe-final.pdf\n- Revisar document DOC-123\n- Falta una dada'
        state, anomalies, documents, original = parse_result(text)
        self.assertEqual(state, 'Pendent')
        self.assertEqual(len(anomalies.splitlines()), 3)
        self.assertIn('informe-final.pdf', documents)
        self.assertIn('DOC-123', documents)
        self.assertIn('3. Sin documento identificable', documents)
        self.assertEqual(original, text)

    def test_inline_list_and_document_link(self):
        result = parse_result('Estado: Pendiente Anomalies: - Revisar solicitud - Falta firma',
                              [{'text': 'solicitud', 'href': 'https://example.com/solicitud.pdf'}])
        self.assertEqual(result[0], 'Pendiente')
        self.assertEqual(len(result[1].splitlines()), 2)
        self.assertIn('https://example.com/solicitud.pdf', result[2])

    def test_absence_is_not_assumed(self):
        self.assertIn('No aparece', parse_result('Estat: Concedit')[1])
        self.assertIn('Sin anomalías', parse_result('Estat: Concedit\nAnomalies: Cap')[1])
        for text in ('', 'Estat: Pendent\nAnomalies:'):
            with self.subTest(text=text), self.assertRaises(ValueError):
                parse_result(text)

    def test_user_example_catalan_document_headings(self):
        text = (
            "En relaió amb l'expedient FACTOR26-9999/2026, una vegada revisada la documentació "
            "presentad en la sol·licitud, s'han notificat les següents anomalies:&#x20;\n\n"
            "- Document normalitzat 1: Sol·licitud: És necessari aportar el document normalitzat 1 "
            "de la sol·licitud de subvenció, degudament emplenat i signat pel sol·licitant "
            "(el document no es troba signat).\n"
            "- Document 2: memòria resum de l'actuació: En el document presentat 2 les dades "
            "del cost de l'actuació no estan emplenades."
        )
        state, anomalies, documents, original = parse_result(text)
        self.assertEqual(state, 'No indicado en el texto')
        self.assertEqual(len(anomalies.splitlines()), 2)
        self.assertNotIn('&#x20;', anomalies)
        self.assertIn('el document no es troba signat', anomalies)
        self.assertIn("les dades del cost de l'actuació no estan emplenades", anomalies)
        self.assertEqual(documents, "1. Document normalitzat 1: Sol·licitud\n2. Document 2: memòria resum de l'actuació")
        self.assertEqual(original, text)

    def test_export_preserves_reference_and_literal_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'results.xlsx'
            parsed = parse_result('Estat: Pendent\nAnomalies: - Revisar document DOC-123')
            save_results(path, [('00123456A', '0001', *parsed, '', 2)])
            book = load_workbook(path)
            try:
                sheet = book.active
                self.assertEqual(sheet['A2'].value, '00123456A')
                self.assertIn('DOC-123', sheet['E2'].value)
                self.assertEqual(sheet['F2'].value, 2)
                self.assertEqual(sheet.max_column, 6)
                self.assertEqual(sheet['C2'].value, 'Anomalías')
                self.assertEqual(len(sheet.conditional_formatting), 1)
                rules = list(sheet.conditional_formatting)[0]
                self.assertEqual(len(sheet.conditional_formatting[rules]), 3)
                self.assertNotIn('Texto original', [cell.value for cell in sheet[1]])
            finally:
                book.close()


if __name__ == '__main__':
    unittest.main()
